"""Audit generated template workbooks against their OCR evidence ledger."""

from __future__ import annotations

import json
from collections import Counter
from math import isclose
from pathlib import Path
import re

from openpyxl import load_workbook

from demo.adapters.audit_intake import _ocr_number, tables_from_pages
from demo.adapters.template_workbook_mapper import (
    _key,
    _page_unit_to_yuan,
    _source_rows,
    _unit_of_sheet,
)
from demo.run_material_intake import _sha256


def _formula_count(workbook) -> int:
    return sum(
        1
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )


def _formula_map(workbook) -> dict[tuple[str, str], str]:
    return {
        (sheet.title, cell.coordinate): cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    }


def _forbidden_cells(workbook) -> list[dict[str, str]]:
    forbidden_company_tokens = ("通富热处理", "通富昆山")
    formula_error_tokens = ("#NAME?", "#REF!", "#VALUE!", "#DIV/0!", "_xlfn.")
    found: list[dict[str, str]] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str):
                    continue
                if any(token in value for token in forbidden_company_tokens):
                    found.append({"cell": f"{sheet.title}!{cell.coordinate}", "issue": "残留通富样例主体文本", "value": value})
                if any(token in value for token in formula_error_tokens):
                    found.append({"cell": f"{sheet.title}!{cell.coordinate}", "issue": "公式或单元格含错误标记", "value": value})
    return found


def verify(root: Path, output_dir: Path, template_dir: Path) -> dict[str, object]:
    manifest = json.loads((output_dir / "生成清单.json").read_text(encoding="utf-8"))
    asset_template = load_workbook(template_dir / "上报表文件_通富昆山_已处理.xlsx", read_only=True, data_only=False)
    income_template = load_workbook(template_dir / "通富热处理（昆山）有限公司-收益法-20250630.xlsx", read_only=True, data_only=False)
    result: dict[str, object] = {
        "projects": [],
        "issues": [],
        "material_count": len(manifest.get("all_materials", [])),
        "materials_with_ocr_pages": sum(1 for item in manifest.get("all_materials", []) if item.get("page_count", 0) > 0),
        "materials_with_issues": sum(1 for item in manifest.get("all_materials", []) if item.get("issues")),
    }
    issues: list[dict[str, object]] = result["issues"]  # type: ignore[assignment]
    for document in manifest["documents"]:
        source_path = root / document["source_file"]
        pages = json.loads((output_dir / "ocr_cache" / f"{_sha256(source_path)}.json").read_text(encoding="utf-8"))
        tables = {(table.page_number, table.table_id): table for table in tables_from_pages(source_path.name, pages)}
        units = _page_unit_to_yuan(pages)
        source_rows = _source_rows(source_path.name, pages)
        extracted_tables = list(tables_from_pages(source_path.name, pages))
        source_locations = {
            (item.table.page_number, item.table.table_id, item.row, column): item
            for candidates in source_rows.values()
            for item in candidates
            for column, _value in item.values
        }
        paths = {name: output_dir / value for name, value in document["workbooks"].items()}
        trace = load_workbook(paths["trace_workbook"], read_only=True, data_only=True)
        mappings = list(trace["AI映射规则"].iter_rows(min_row=2, values_only=True))
        workbooks = {
            "资产基础法_资产清查.xlsx": load_workbook(paths["reporting_workbook"], read_only=True, data_only=False),
            "收益法_市场法.xlsx": load_workbook(paths["income_workbook"], read_only=True, data_only=False),
        }
        expected_templates = {
            "资产基础法_资产清查.xlsx": asset_template,
            "收益法_市场法.xlsx": income_template,
        }
        checked, filled, unfilled = 0, 0, 0
        semantic_checked = 0
        unfilled_reasons: Counter[str] = Counter()
        for target_file, workbook in workbooks.items():
            template = expected_templates[target_file]
            if workbook.sheetnames != template.sheetnames:
                issues.append({"source_file": document["source_file"], "issue": f"{target_file}: Sheet 结构不一致"})
            generated_formulas = _formula_map(workbook)
            template_formulas = _formula_map(template)
            if generated_formulas != template_formulas:
                issues.append({
                    "source_file": document["source_file"],
                    "issue": f"{target_file}: 公式坐标或公式文本与模板不一致",
                    "generated_formula_count": len(generated_formulas),
                    "template_formula_count": len(template_formulas),
                })
            for forbidden in _forbidden_cells(workbook):
                issues.append({"source_file": document["source_file"], "target_file": target_file, **forbidden})
        required_trace_sheets = {"AI映射规则", "AI公式规则", "模板字段注册表", "附注及明细候选", "审计原始索引", "企业资料索引", "项目主体信息"}
        missing_trace_sheets = sorted(required_trace_sheets.difference(trace.sheetnames))
        if missing_trace_sheets:
            issues.append({"source_file": document["source_file"], "issue": "追溯工作簿缺少规则工作表", "missing": missing_trace_sheets})
        for mapping in mappings:
            target_file, sheet_name, cell, _field, _source_file, page, table_id, source_rc, _method, _calculation, status = mapping
            if status == "未填":
                unfilled += 1
                unfilled_reasons[str(_calculation or _method or "未说明")] += 1
                continue
            if status != "已填":
                continue
            filled += 1
            # This verifier's scalar equality check applies only to statutory
            # statement mappings.  A note-derived mapping may also cite one
            # exact ``row,column`` pair, but can legitimately apply unit,
            # sign, percentage, category aggregation or reconciliation rules;
            # treating it as a raw scalar copy creates false failures.
            if (
                _method != "源值直取+单位换算"
                or not isinstance(source_rc, str)
                or not re.fullmatch(r"\d+,\d+", source_rc.strip())
            ):
                continue
            source_row, source_column = (int(part) for part in source_rc.split(","))
            table = tables.get((int(page), str(table_id)))
            if table is None or source_row > len(table.matrix) or source_column > len(table.matrix[source_row - 1]):
                issues.append({"source_file": document["source_file"], "cell": f"{sheet_name}!{cell}", "issue": "来源定位失效"})
                continue
            raw = _ocr_number(table.matrix[source_row - 1][source_column - 1])
            actual = workbooks[target_file][sheet_name][cell].value
            # The income template stores audited historical statements in
            # 万元 even though the unit label is outside the compact scan used
            # by the generic sheet-unit helper.
            target_unit = 10_000.0 if target_file == "收益法_市场法.xlsx" and sheet_name in {"历资表", "历利表", "历现表"} else _unit_of_sheet(workbooks[target_file][sheet_name])
            source_unit = units.get(int(page))
            expected = float(raw) * source_unit / target_unit if isinstance(raw, (int, float)) and source_unit is not None else None
            checked += 1
            if not isinstance(actual, (int, float)) or expected is None or not isclose(float(actual), expected, rel_tol=1e-9, abs_tol=1e-7):
                issues.append({"source_file": document["source_file"], "cell": f"{sheet_name}!{cell}", "actual": actual, "expected": expected, "issue": "取数或单位换算不一致"})
            source = source_locations.get((int(page), str(table_id), source_row, source_column))
            if source is None:
                issues.append({"source_file": document["source_file"], "cell": f"{sheet_name}!{cell}", "issue": "来源列未落在标准科目段内"})
                continue
            semantic_checked += 1
            if _key(_field) != _key(source.label):
                issues.append({
                    "source_file": document["source_file"],
                    "cell": f"{sheet_name}!{cell}",
                    "target_field": _field,
                    "source_label": source.label,
                    "issue": "目标科目与来源科目语义不一致",
                })
            if source.unit_to_yuan is None:
                issues.append({"source_file": document["source_file"], "cell": f"{sheet_name}!{cell}", "issue": "已填值缺少明确来源单位"})
            if source_column not in dict(source.periods):
                issues.append({"source_file": document["source_file"], "cell": f"{sheet_name}!{cell}", "issue": "已填值缺少明确来源期间"})
        registry = trace["模板字段注册表"] if "模板字段注册表" in trace.sheetnames else None
        registry_rows = max((registry.max_row - 1), 0) if registry is not None else 0
        registry_missing = 0
        if registry is not None:
            for row in registry.iter_rows(min_row=2, values_only=True):
                if row[5] in {"未填", "待补充"}:
                    registry_missing += 1
                    if not row[7]:
                        issues.append({"source_file": document["source_file"], "cell": f"{row[1]}!{row[2]}", "issue": "缺失字段未注明补充资料"})
        result["projects"].append({
            "source_file": document["source_file"],
            "filled": filled,
            "unfilled": unfilled,
            "checked": checked,
            "semantic_checked": semantic_checked,
            "template_registry_rows": registry_rows,
            "template_registry_missing": registry_missing,
            "source_table_categories": dict(Counter(table.category for table in extracted_tables)),
            "source_standard_rows": sum(len(items) for items in source_rows.values()),
            "source_rows_with_known_unit": sum(1 for items in source_rows.values() for item in items if item.unit_to_yuan is not None),
            "source_rows_with_complete_periods": sum(
                1
                for items in source_rows.values()
                for item in items
                if item.values and all(column in dict(item.periods) for column, _value in item.values)
            ),
            "unfilled_reasons": dict(unfilled_reasons),
        })  # type: ignore[index]
    result["summary"] = {
        "project_count": len(result["projects"]),  # type: ignore[arg-type]
        "filled": sum(item["filled"] for item in result["projects"]),  # type: ignore[index]
        "unfilled": sum(item["unfilled"] for item in result["projects"]),  # type: ignore[index]
        "checked": sum(item["checked"] for item in result["projects"]),  # type: ignore[index]
        "semantic_checked": sum(item["semantic_checked"] for item in result["projects"]),  # type: ignore[index]
        "issue_count": len(issues),
    }
    return result


if __name__ == "__main__":
    checked = verify(
        Path(r"D:\资产评估\00审计报告-测试"),
        Path(r"D:\AssetAppraisal\outputs\audit-material-intake"),
        Path(r"D:\AssetAppraisal\资产评估工作流"),
    )
    report_path = Path(r"D:\AssetAppraisal\outputs\audit-material-intake\workbooks_generic_safe_v6\全量核验报告.json")
    report_path.write_text(json.dumps(checked, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(checked, ensure_ascii=False, indent=2))
