from openpyxl import load_workbook

from demo.adapters.audit_intake import ROLE_INCOME, ROLE_REPORTING, build_role_workbooks, needs_ocr, tables_from_pages
from demo.adapters.rule_graph_builder import build_rule_graph_bundle
from demo.run_material_intake import _structured_pages


def _pages():
    return [
        {
            "page_number": 5,
            "blocks": [{"text": "资产负债表 单位：元"}],
            "tables": [{"table_id": "p5-t1", "cells": [
                {"row": 1, "column": 1, "text": "项目"},
                {"row": 1, "column": 2, "text": "2025年6月30日"},
                {"row": 2, "column": 1, "text": "资产总计"},
                {"row": 2, "column": 2, "text": "1,234.50"},
                {"row": 3, "column": 1, "text": "负债合计"},
                {"row": 3, "column": 2, "text": "234.50"},
            ]}],
        },
        {
            "page_number": 6,
            "blocks": [{"text": "利润表 单位：元"}],
            "tables": [{"table_id": "p6-t1", "cells": [
                {"row": 1, "column": 1, "text": "项目"},
                {"row": 1, "column": 2, "text": "2025年1-6月"},
                {"row": 1, "column": 3, "text": "2024年1-6月"},
                {"row": 2, "column": 1, "text": "净利润"},
                {"row": 2, "column": 2, "text": "100"},
                {"row": 2, "column": 3, "text": "90"},
                {"row": 3, "column": 1, "text": "营业收入"},
                {"row": 3, "column": 2, "text": "1,000"},
                {"row": 3, "column": 3, "text": "900"},
            ]}],
        },
    ]


def test_classifies_financial_tables_by_content_not_file_name():
    tables = tables_from_pages("任意文件.pdf", _pages())
    assert tables[0].roles == (ROLE_REPORTING,)
    assert tables[1].roles == (ROLE_INCOME,)


def test_does_not_classify_an_audit_procedure_table_from_page_text_alone():
    pages = [{
        "page_number": 1,
        "blocks": [{"text": "资产负债表"}],
        "tables": [{"table_id": "audit-matter", "cells": [
            {"row": 1, "column": 1, "text": "关键审计事项"},
            {"row": 1, "column": 2, "text": "审计应对"},
            {"row": 2, "column": 1, "text": "收入确认"},
        ]}],
    }]
    assert tables_from_pages("任意文件.pdf", pages)[0].roles == ()


def test_only_requests_ocr_when_local_pdf_has_no_text_or_table():
    assert needs_ocr([{ "blocks": [{"text": "审计报告"}], "tables": [] }]) is False
    assert needs_ocr([{ "blocks": [{"text": ""}], "tables": [] }]) is True


def test_builds_two_auditable_role_workbooks(tmp_path):
    paths = build_role_workbooks(output_dir=tmp_path, document_name="任意文件.pdf", pages=_pages())
    reporting = load_workbook(paths[ROLE_REPORTING], data_only=False)
    income = load_workbook(paths[ROLE_INCOME], data_only=False)
    assert any(name.startswith("资产负债表") for name in reporting.sheetnames)
    assert any(name.startswith("利润表") for name in income.sheetnames)
    assert reporting["映射规则"].max_row > 1
    assert income["映射规则"].max_row > 1
    assert (tmp_path / "映射关系图.mmd").is_file()


def test_adds_an_auditable_accounting_identity_when_statement_rows_allow_it(tmp_path):
    pages = [{
        "page_number": 1,
        "blocks": [{"text": "资产负债表"}],
        "tables": [{"table_id": "p1-t1", "cells": [
            {"row": 1, "column": 1, "text": "项目"},
            {"row": 1, "column": 2, "text": "2025年6月30日"},
            {"row": 2, "column": 1, "text": "流动资产合计"},
            {"row": 2, "column": 2, "text": "100"},
            {"row": 3, "column": 1, "text": "非流动资产合计"},
            {"row": 3, "column": 2, "text": "200"},
            {"row": 4, "column": 1, "text": "资产总计"},
            {"row": 4, "column": 2, "text": "300"},
        ]}],
    }]
    path = build_role_workbooks(output_dir=tmp_path, document_name="任意文件.pdf", pages=pages)[ROLE_REPORTING]
    workbook = load_workbook(path, data_only=False)
    assert "核验计算" in workbook.sheetnames
    assert "资产负债表_1" in workbook["核验计算"]["C2"].value


def test_rule_graph_browser_uses_role_colours_and_shows_source_evidence(tmp_path):
    workbook_path = tmp_path / "资产法.xlsx"
    workbook = load_workbook(build_role_workbooks(
        output_dir=tmp_path / "workbooks", document_name="审计报告.pdf", pages=_pages(),
    )[ROLE_REPORTING])
    sheet_name = next(name for name in workbook.sheetnames if name.startswith("资产负债表"))
    workbook[sheet_name]["B2"] = "=B3"
    workbook.save(workbook_path)
    trace = tmp_path / "trace.xlsx"
    trace_book = load_workbook(workbook_path)
    trace_sheet = trace_book.create_sheet("AI映射规则")
    trace_sheet.append(["目标文件", "目标工作表", "目标单元格", "模板字段", "来源文件", "来源页", "来源表", "来源行", "取数方式", "计算/处理", "状态"])
    trace_sheet.append(["资产法", sheet_name, "B3", "资产总计", "审计报告.pdf", 5, "资产负债表", "第2行", "直接取数", "单位按原报表", "已填"])
    trace_book.save(trace)

    build_rule_graph_bundle(
        output_dir=tmp_path / "graph",
        workbook_paths={"资产法": workbook_path},
        trace_path=trace,
    )

    browser = (tmp_path / "graph" / "逐单元格规则图浏览器.html").read_text(encoding="utf-8")
    assert ".source_evidence rect{fill:#fff8cf" in browser
    assert "来源证据" in browser
    assert "重点规则节点" in browser
    assert "evidence_mapping" in browser


def test_material_intake_reads_native_docx_and_xlsx_content(tmp_path):
    document_path = tmp_path / "企业介绍.docx"
    document = __import__("docx").Document()
    document.add_paragraph("某某公司企业介绍")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "主营业务"
    table.cell(0, 1).text = "热处理"
    document.save(document_path)
    assert "某某公司企业介绍" in _structured_pages(document_path)[0]["blocks"][0]["text"]

    workbook_path = tmp_path / "审计附表.xlsx"
    workbook = load_workbook(build_role_workbooks(
        output_dir=tmp_path / "native-xlsx", document_name="审计报告.pdf", pages=_pages(),
    )[ROLE_REPORTING])
    worksheet = workbook.active
    worksheet["A1"] = "资产负债表"
    worksheet["B1"] = "资产总计"
    workbook.save(workbook_path)
    pages = _structured_pages(workbook_path)
    assert any(cell["text"] == "资产负债表" for page in pages for table in page["tables"] for cell in table["cells"])
