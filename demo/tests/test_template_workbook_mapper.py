from demo.adapters.audit_intake import ExtractedTable
from openpyxl import Workbook

from demo.adapters.template_workbook_mapper import (
    SourceRow,
    _apply_project_identity,
    _clear_asset_sample_data,
    _asset_target_accepts_statement_value,
    _find_source,
)


def test_find_source_falls_back_to_unknown_scope_only_when_preferred_scope_missing():
    table = ExtractedTable(
        source_file="审计报告.pdf",
        page_number=1,
        table_id="t1",
        category="利润表",
        roles=("利润表",),
        matrix=[],
        period_headers={},
        confidence="high",
    )
    source_rows = {
        "营业收入": [
            SourceRow(table, 2, "营业收入", ((2, 1000.0),), ((2, "2025年度"),), 1.0, "未识别"),
        ]
    }

    source = _find_source("营业收入", source_rows, preferred_scope="母公司")

    assert source is not None
    assert source.scope == "未识别"


def test_find_source_uses_requested_scope_when_both_explicit_scopes_exist():
    table = ExtractedTable(
        source_file="审计报告.pdf",
        page_number=1,
        table_id="t1",
        category="利润表",
        roles=("利润表",),
        matrix=[],
        period_headers={},
        confidence="high",
    )
    source_rows = {
        "营业收入": [
            SourceRow(table, 2, "营业收入", ((2, 1000.0),), ((2, "2025年度"),), 1.0, "母公司"),
            SourceRow(table, 3, "营业收入", ((2, 2000.0),), ((2, "2025年度"),), 1.0, "合并"),
        ]
    }

    source = _find_source("营业收入", source_rows, preferred_scope="母公司")

    assert source is not None
    assert source.scope == "母公司"


def test_asset_statement_value_is_not_written_into_original_cost_column():
    workbook = Workbook()
    sheet = workbook.active
    sheet["D7"] = "账面原值"
    sheet["C8"] = "固定资产合计"

    assert not _asset_target_accepts_statement_value(sheet, 8, 4)


def test_clear_asset_details_includes_rows_beyond_original_capacity():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "表4-6-6"
    sheet["C7"] = "设备名称"
    sheet["C49"] = "旧项目设备"
    sheet["J65"] = 12345
    sheet["C66"] = "合计"
    sheet["J66"] = "=SUM(J8:J65)"
    _clear_asset_sample_data(workbook)
    assert sheet["C49"].value is None
    assert sheet["J65"].value is None
    assert sheet["C7"].value == "设备名称"
    assert sheet["C66"].value == "合计"
    assert sheet["J66"].value == "=SUM(J8:J65)"


def test_asset_statement_value_maps_only_to_net_row_of_gross_allowance_net_block():
    workbook = Workbook()
    sheet = workbook.active
    sheet["D7"] = "账面价值"
    sheet["C8"] = "存货合计"
    sheet["C9"] = "减：存货跌价准备"
    sheet["C10"] = "存货净额"

    assert not _asset_target_accepts_statement_value(sheet, 8, 4)
    assert not _asset_target_accepts_statement_value(sheet, 9, 4)
    assert _asset_target_accepts_statement_value(sheet, 10, 4)


def test_project_identity_fills_company_neutral_template_placeholders():
    asset = Workbook()
    asset.active["A1"] = "被评估单位：{{COMPANY_NAME}}"
    asset.active["A2"] = "{{COMPANY_SHORT_NAME}}"
    asset.active["A3"] = "{{APPRAISAL_ORG_NAME}}"
    asset.active["A4"] = "评估基准日: {{VALUATION_BASE_DATE_CN}}"
    income = Workbook()
    income.active.title = "项目信息"
    income.active["A2"] = "评估基准日：{{VALUATION_BASE_DATE}}"

    _apply_project_identity(
        asset,
        income,
        {
            "company_name": "示例科技有限公司",
            "company_short_name": "示例科技",
            "appraisal_organization_name": "示例评估机构",
            "legal_representative": "张三",
        },
        "2026-06-30",
    )

    assert asset.active["A1"].value == "被评估单位：示例科技有限公司"
    assert asset.active["A2"].value == "示例科技"
    assert asset.active["A3"].value == "示例评估机构"
    assert asset.active["A4"].value == "评估基准日:  2026年 06月 30日"
    assert income["项目信息"]["A2"].value == "评估基准日：2026-06-30"
    assert income["项目信息"]["B5"].value == "示例科技有限公司"
