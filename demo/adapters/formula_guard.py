"""Guard formula display errors that are caused by intentionally blank inputs."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


GUARDABLE_ERRORS = {"#DIV/0!", "#VALUE!", "#N/A", "#NUM!", "#NULL!"}
STRUCTURAL_ERRORS = {"#NAME?", "#REF!"}


def guard_cached_formula_errors(path: Path) -> dict[str, Any]:
    """Wrap only formulas whose Excel-cached result is a guardable error.

    The workbook must first have been calculated by Excel/LibreOffice.  This
    routine never masks #NAME? or #REF!: those indicate broken formulas and
    must remain hard QA failures.
    """
    formulas = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    guarded: list[dict[str, str]] = []
    structural: list[dict[str, str]] = []
    for cached_sheet in cached.worksheets:
        formula_sheet = formulas[cached_sheet.title]
        for cached_cell in cached_sheet._cells.values():
            value = cached_cell.value
            if value not in GUARDABLE_ERRORS | STRUCTURAL_ERRORS:
                continue
            formula_cell = formula_sheet[cached_cell.coordinate]
            formula = formula_cell.value
            item = {"sheet": cached_sheet.title, "cell": cached_cell.coordinate, "error": str(value), "formula": str(formula)}
            if value in STRUCTURAL_ERRORS or not (isinstance(formula, str) and formula.startswith("=")):
                structural.append(item)
                continue
            if not formula.upper().startswith("=IFERROR("):
                formula_cell.value = f'=IFERROR({formula[1:]},"")'
                guarded.append(item)
    cached.close()
    if guarded:
        formulas.calculation.calcMode = "auto"
        formulas.calculation.fullCalcOnLoad = True
        formulas.calculation.forceFullCalc = True
        formulas.save(path)
    formulas.close()
    return {"guarded_count": len(guarded), "guarded": guarded, "structural_errors": structural}
