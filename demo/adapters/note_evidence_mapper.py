"""Map audited financial-statement notes into the two appraisal templates.

The statutory statements prove totals.  Their notes often prove a second,
more detailed layer (inventory classes, fixed-asset classes, expense lines,
etc.).  This module keeps that distinction explicit:

* only the selected reporting scope is used;
* every written cell carries a page/table/row locator;
* note lines are never treated as forecasts or appraisal conclusions;
* when a note says ``其中主要为`` and omits part of a total, the unexplained
  balance is written as a transparent, calculated residual rather than being
  assigned to a fabricated account.
"""

from __future__ import annotations

from dataclasses import dataclass
import re
from typing import Any, Iterable

from openpyxl.utils import get_column_letter

from .audit_intake import ExtractedTable, _ocr_number, tables_from_pages


NOTE_TOPICS = (
    "货币资金", "应收账款", "预付款项", "其他应收款", "存货", "其他流动资产",
    "长期股权投资", "固定资产", "在建工程", "使用权资产", "无形资产",
    "长期待摊费用", "短期借款", "应付账款", "应付职工薪酬", "应交税费",
    "其他应付款", "一年内到期的非流动负债", "长期借款", "租赁负债",
    "实收资本", "资本公积", "未分配利润", "营业收入和营业成本",
    "税金及附加", "销售费用", "管理费用", "研发费用", "财务费用",
    "其他收益", "信用减值损失", "资产减值损失", "资产处置收益",
    "营业外收入", "营业外支出", "所得税费用",
)


@dataclass(frozen=True)
class NoteTable:
    table: ExtractedTable
    scope: str
    topic: str


@dataclass(frozen=True)
class NoteValue:
    label: str
    period: int
    value: float
    table: ExtractedTable
    row: int
    column: int


def _compact(value: Any) -> str:
    text = str(value or "").strip()
    text = re.sub(r"^(?:其中主要为|其中|加|减)[：:]?", "", text)
    text = re.sub(r"[（(].*?[）)]", "", text)
    text = re.sub(r"[\s：:、，,。·—\-]", "", text)
    return text


def _heading_topic(text: str) -> str | None:
    cleaned = re.sub(r"^[（(][一二三四五六七八九十百]+[）)]\s*", "", text.strip())
    cleaned = re.sub(r"^\d+[、.]\s*", "", cleaned)
    matches = [topic for topic in NOTE_TOPICS if topic in cleaned]
    return max(matches, key=len) if matches else None


def note_tables_from_pages(document_name: str, pages: list[dict[str, Any]]) -> list[NoteTable]:
    """Attach a reporting scope and note heading to every OCR table.

    DocMind preserves table order but not table bounding boxes.  A continuation
    table therefore inherits the previous heading.  When a page has more
    tables than note headings, the leading tables are continuations and the
    trailing tables pair with the headings in display order.
    """
    tables_by_page: dict[int, list[ExtractedTable]] = {}
    for table in tables_from_pages(document_name, pages):
        tables_by_page.setdefault(table.page_number, []).append(table)

    scope = "未识别"
    current_topic = ""
    result: list[NoteTable] = []
    for page in pages:
        page_number = int(page.get("page_number") or 0)
        block_texts = [str(block.get("text") or "").strip() for block in page.get("blocks", [])]
        joined = " ".join(block_texts)
        if "母公司财务报表项目注释" in joined:
            scope = "母公司"
        elif "合并财务报表项目注释" in joined:
            scope = "合并"
        headings: list[str] = []
        for text in block_texts:
            topic = _heading_topic(text)
            if topic and (not headings or headings[-1] != topic):
                headings.append(topic)
        page_tables = tables_by_page.get(page_number, [])
        if not page_tables:
            if headings:
                current_topic = headings[-1]
            continue
        continuation_count = max(0, len(page_tables) - len(headings))
        for index, table in enumerate(page_tables):
            if index < continuation_count:
                topic = current_topic
            else:
                heading_index = index - continuation_count
                topic = headings[heading_index] if heading_index < len(headings) else current_topic
            result.append(NoteTable(table=table, scope=scope, topic=topic or "未分类附注"))
        if headings:
            current_topic = headings[-1]
    return result


def _year_columns(matrix: list[list[str]]) -> dict[int, int]:
    result: dict[int, int] = {}
    for row in matrix[:3]:
        carried: int | None = None
        for column, raw in enumerate(row, 1):
            match = re.search(r"((?:19|20)\d{2})\s*年", str(raw or ""))
            if match:
                carried = int(match.group(1))
            if carried is not None:
                result.setdefault(column, carried)
    return result


def _simple_note_values(note: NoteTable) -> list[NoteValue]:
    matrix = note.table.matrix
    years = _year_columns(matrix)
    result: list[NoteValue] = []
    if not years:
        return result
    for row_number, row in enumerate(matrix, 1):
        label = str(row[0] if row else "").strip()
        if not label or re.search(r"项目|账龄|资产类别|被投资单位|税费项目", label):
            continue
        for column, year in years.items():
            raw = row[column - 1] if len(row) >= column else ""
            value = _ocr_number(raw)
            if isinstance(value, (int, float)):
                result.append(NoteValue(label, year, float(value), note.table, row_number, column))
    return result


def _selected(notes: Iterable[NoteTable], topic: str, scope: str) -> list[NoteTable]:
    return [note for note in notes if note.topic == topic and note.scope == scope]


def _source_fields(values: Iterable[NoteValue]) -> tuple[str, str, str, str]:
    values = list(values)
    files = sorted({value.table.source_file for value in values})
    pages = sorted({value.table.page_number for value in values})
    tables = sorted({value.table.table_id for value in values})
    rows = sorted({f"{value.row},{value.column}" for value in values})
    return (
        "；".join(files),
        ",".join(map(str, pages)),
        "；".join(tables),
        "；".join(rows),
    )


def _append_mapping(
    mappings: list[list[Any]], sheet, cell, field: str, values: Iterable[NoteValue],
    method: str, transform: str, status: str = "已填",
) -> None:
    source_file, pages, tables, rows = _source_fields(values)
    mappings.append([
        sheet.title, cell.coordinate, field, source_file, pages, tables, rows,
        method, transform, status,
    ])


def _target_year_columns(sheet) -> dict[int, int]:
    result: dict[int, int] = {}
    for row in range(1, min(sheet.max_row, 8) + 1):
        for column in range(1, min(sheet.max_column, 20) + 1):
            match = re.search(r"((?:19|20)\d{2})\s*年", str(sheet.cell(row, column).value or ""))
            if match:
                result.setdefault(int(match.group(1)), column)
    return result


def _set_income_value(sheet, cell, amount_yuan: float, evidence: list[NoteValue], mappings, field: str, method: str = "附注直取") -> None:
    if isinstance(cell.value, str) and cell.value.startswith("="):
        return
    cell.value = amount_yuan / 10_000.0
    _append_mapping(mappings, sheet, cell, field, evidence, method, "人民币元÷10,000写入万元；母公司口径")


def _income_statement_split(income, notes: list[NoteTable], mappings, scope: str) -> None:
    candidates = _selected(notes, "营业收入和营业成本", scope)
    if not candidates:
        return
    sheet = income["历利表"]
    target = {
        ("主营业务", "收入"): 8,
        ("其他业务", "收入"): 9,
        ("主营业务", "成本"): 15,
        ("其他业务", "成本"): 16,
    }
    for note in candidates:
        matrix = note.table.matrix
        if len(matrix) < 3 or not any("收入" in str(value) for value in matrix[1]):
            continue
        years = _year_columns(matrix)
        measures: dict[int, str] = {}
        for column, raw in enumerate(matrix[1], 1):
            text = str(raw or "").strip()
            if text in {"收入", "成本"}:
                measures[column] = text
        for row_number, row in enumerate(matrix[2:], 3):
            business = str(row[0] if row else "").strip()
            if business not in {"主营业务", "其他业务"}:
                continue
            for column, measure in measures.items():
                year = years.get(column)
                if year not in {2023, 2024, 2025}:
                    continue
                value = _ocr_number(row[column - 1] if len(row) >= column else "")
                if not isinstance(value, (int, float)):
                    continue
                target_column = {2023: 3, 2024: 4, 2025: 5}[year]
                cell = sheet.cell(target[(business, measure)], target_column)
                evidence = [NoteValue(business, year, float(value), note.table, row_number, column)]
                _set_income_value(sheet, cell, float(value), evidence, mappings, f"{business}{measure}")


INCOME_NOTE_ROWS = {
    "税金及附加": 24, "销售费用": 25, "管理费用": 26, "研发费用": 27,
    "财务费用": 28, "其他收益": 31, "资产减值损失": 39,
    "资产处置收益": 40, "营业外收入": 42, "营业外支出": 43,
}


def _note_topic_totals(notes: list[NoteTable], topic: str, scope: str) -> dict[int, tuple[float, list[NoteValue]]]:
    values = [value for note in _selected(notes, topic, scope) for value in _simple_note_values(note)]
    explicit: dict[int, list[NoteValue]] = {}
    detail: dict[int, list[NoteValue]] = {}
    for value in values:
        if "合计" in _compact(value.label) or "总计" in _compact(value.label):
            explicit.setdefault(value.period, []).append(value)
        else:
            detail.setdefault(value.period, []).append(value)
    result: dict[int, tuple[float, list[NoteValue]]] = {}
    for year in set(explicit) | set(detail):
        evidence = explicit.get(year) or detail.get(year) or []
        result[year] = (sum(item.value for item in evidence), evidence)
    return result


def _fill_income_totals_from_notes(income, notes, mappings, scope: str) -> dict[int, dict[int, tuple[float, list[NoteValue]]]]:
    """Use note totals only where the statutory statement OCR is unreadable."""
    sheet = income["历利表"]
    totals_by_row: dict[int, dict[int, tuple[float, list[NoteValue]]]] = {}
    for topic, row in INCOME_NOTE_ROWS.items():
        totals = _note_topic_totals(notes, topic, scope)
        totals_by_row[row] = totals
        for year, (amount, evidence) in totals.items():
            column = {2023: 3, 2024: 4, 2025: 5}.get(year)
            if not column or not evidence or sheet.cell(row, column).value not in (None, ""):
                continue
            _set_income_value(
                sheet, sheet.cell(row, column), amount, evidence, mappings, topic,
                "审计主表OCR不可读时，以同口径附注合计补足",
            )
    return totals_by_row


def _statement_amount(source_rows: dict[str, list[Any]], key: str, scope: str, year: int) -> float | None:
    candidates = [item for item in source_rows.get(key, []) if getattr(item, "scope", None) == scope]
    if not candidates:
        return None
    source = max(candidates, key=lambda item: len(item.values))
    periods = dict(source.periods)
    for column, value in source.values:
        if str(periods.get(column, "")).startswith(str(year)) and source.unit_to_yuan:
            return float(value) * float(source.unit_to_yuan) / 10_000.0
    return None


def _reconcile_operating_profit(
    income, totals_by_row: dict[int, dict[int, tuple[float, list[NoteValue]]]],
    mappings, source_rows: dict[str, list[Any]], scope: str,
) -> None:
    """Correct isolated OCR digits only when note evidence restores the audited equation."""
    sheet = income["历利表"]
    expense_rows = {24, 25, 26, 27, 28}
    candidate_rows = {24, 25, 26, 27, 28, 31, 39, 40}
    for year, column in ((2023, 3), (2024, 4), (2025, 5)):
        audited = _statement_amount(source_rows, "营业利润", scope, year)
        if audited is None:
            continue
        numeric = lambda row: float(sheet.cell(row, column).value or 0.0) if isinstance(sheet.cell(row, column).value, (int, float)) else 0.0
        revenue = numeric(8) + numeric(9) + sum(numeric(row) for row in range(10, 13))
        cost = numeric(15) + numeric(16) + sum(numeric(row) for row in range(17, 29))
        calculated = revenue - cost + numeric(31) + numeric(32) + sum(numeric(row) for row in range(35, 41))
        needed = audited - calculated
        candidates: list[tuple[int, float, float, list[NoteValue]]] = []
        for row in candidate_rows:
            current = sheet.cell(row, column).value
            note = totals_by_row.get(row, {}).get(year)
            if not isinstance(current, (int, float)) or not note:
                continue
            note_amount = note[0] / 10_000.0
            delta = note_amount - float(current)
            effect = -delta if row in expense_rows else delta
            if abs(effect) > 0.000001:
                candidates.append((row, effect, note_amount, note[1]))
        # Exhaustive subset search is safe here (at most eight audited lines)
        # and prevents choosing a plausible-looking correction by intuition.
        best: tuple[float, list[tuple[int, float, float, list[NoteValue]]]] | None = None
        for mask in range(1, 1 << len(candidates)):
            chosen = [candidate for index, candidate in enumerate(candidates) if mask & (1 << index)]
            residual = abs(needed - sum(item[1] for item in chosen))
            if best is None or residual < best[0]:
                best = (residual, chosen)
        if best is None or best[0] > 0.0001:
            continue
        for row, _effect, note_amount, evidence in best[1]:
            cell = sheet.cell(row, column)
            cell.value = note_amount
            _append_mapping(
                mappings, sheet, cell, sheet.cell(row, 1).value, evidence,
                "审计附注交叉校正主表OCR", "附注金额替换OCR疑似错位字符；替换后营业利润与审计主表勾稽一致", "已填",
            )


EMPLOYEE_LABELS = {
    "工资奖金", "福利费", "社会保险费", "住房公积金", "工会经费", "职工教育经费",
    "残疾人保障基金", "离职补偿金", "员工赔偿款",
}


def _expense_details(
    income, notes: list[NoteTable], mappings, topic: str, sheet_name: str,
    total_row: int, scope: str, statement_row: int,
) -> None:
    candidates = _selected(notes, topic, scope)
    values = [value for note in candidates for value in _simple_note_values(note)]
    if not values:
        return
    sheet = income[sheet_name]
    year_columns = _target_year_columns(sheet)
    detail: dict[str, dict[int, list[NoteValue]]] = {}
    reported: dict[int, list[NoteValue]] = {}
    for value in values:
        key = _compact(value.label)
        if "合计" in key or "总计" in key:
            reported.setdefault(value.period, []).append(value)
            continue
        detail.setdefault(key, {}).setdefault(value.period, []).append(value)

    employee: dict[int, list[NoteValue]] = {}
    depreciation: dict[int, list[NoteValue]] = {}
    amortization: dict[int, list[NoteValue]] = {}
    right_use: dict[int, list[NoteValue]] = {}
    remaining: dict[str, dict[int, list[NoteValue]]] = {}
    for key, periods in detail.items():
        bucket = remaining
        if key in EMPLOYEE_LABELS:
            for year, items in periods.items(): employee.setdefault(year, []).extend(items)
            continue
        if "使用权资产" in key and "折旧" in key:
            for year, items in periods.items(): right_use.setdefault(year, []).extend(items)
            continue
        if "折旧" in key:
            for year, items in periods.items(): depreciation.setdefault(year, []).extend(items)
            continue
        if "摊销" in key:
            for year, items in periods.items(): amortization.setdefault(year, []).extend(items)
            continue
        bucket[key] = periods

    rows_by_key = {_compact(sheet.cell(row, 1).value): row for row in range(6, total_row) if sheet.cell(row, 1).value}
    alias_rows = {
        "人工": rows_by_key.get("人工"), "折旧": rows_by_key.get("折旧"),
        "摊销": rows_by_key.get("摊销"), "使用权资产折旧": rows_by_key.get("使用权资产折旧"),
    }
    used_rows = {row for row in rows_by_key.values() if row}

    def write_group(name: str, grouped: dict[int, list[NoteValue]]) -> None:
        row = alias_rows.get(name)
        if not row:
            return
        for year, items in grouped.items():
            column = year_columns.get(year)
            if column:
                _set_income_value(sheet, sheet.cell(row, column), sum(item.value for item in items), items, mappings, name, "附注明细汇总")

    write_group("人工", employee)
    write_group("折旧", depreciation)
    write_group("摊销", amortization)
    write_group("使用权资产折旧", right_use)

    aliases = {
        "房租费": "房屋租赁费", "运输快递费": "运输费",
        "利息费用": "利息支出",
    }
    for key, periods in remaining.items():
        target_key = _compact(aliases.get(key, key))
        row = rows_by_key.get(target_key)
        if row is None:
            row = next((candidate for candidate in range(6, total_row) if candidate not in used_rows and sheet.cell(candidate, 1).value in (None, "")), None)
            if row is None:
                continue
            original_label = next(iter(periods.values()))[0].label
            sheet.cell(row, 1).value = re.sub(r"^其中主要为[：:]?", "", original_label)
            used_rows.add(row)
            rows_by_key[target_key] = row
        factor = -1.0 if topic == "财务费用" and key == "利息收入" else 1.0
        for year, items in periods.items():
            column = year_columns.get(year)
            if column:
                _set_income_value(sheet, sheet.cell(row, column), sum(item.value for item in items) * factor, items, mappings, sheet.cell(row, 1).value, "附注明细直取/同名合并+符号转换" if factor < 0 else "附注明细直取/同名合并")

    # Reconcile an explicitly reported note total.  A residual is a disclosed
    # calculation, not a fabricated expense account.
    residual_row: int | None = None
    for year, totals in reported.items():
        column = year_columns.get(year)
        if not column:
            continue
        main_column = {2023: 3, 2024: 4, 2025: 5}.get(year)
        main_value = income["历利表"].cell(statement_row, main_column).value if main_column else None
        reported_total = float(main_value) * 10_000.0 if isinstance(main_value, (int, float)) else totals[-1].value
        written = 0.0
        for row in range(6, total_row):
            value = sheet.cell(row, column).value
            if isinstance(value, (int, float)):
                written += float(value) * 10_000.0
        residual = reported_total - written
        if abs(residual) <= 0.01:
            continue
        if residual_row is None:
            residual_row = next((candidate for candidate in range(total_row - 1, 5, -1) if sheet.cell(candidate, 1).value in (None, "")), None)
            if residual_row is None:
                continue
            sheet.cell(residual_row, 1).value = "审计附注未列示明细差额（倒挤）"
        cell = sheet.cell(residual_row, column)
        _set_income_value(sheet, cell, residual, totals, mappings, sheet.cell(residual_row, 1).value, "审计主表审定总额减已披露附注明细")


def _small_detail_sheet(
    income, notes, mappings, topic: str, sheet_name: str, scope: str,
    statement_row: int, sign_rules: dict[str, float] | None = None,
) -> None:
    candidates = _selected(notes, topic, scope)
    values = [value for note in candidates for value in _simple_note_values(note)]
    if not values:
        return
    sheet = income[sheet_name]
    year_columns = _target_year_columns(sheet)
    grouped: dict[str, dict[int, list[NoteValue]]] = {}
    for value in values:
        key = _compact(value.label)
        if "合计" in key or "总计" in key:
            continue
        grouped.setdefault(key, {}).setdefault(value.period, []).append(value)
    last_row = 5
    for offset, (key, periods) in enumerate(grouped.items()):
        row = 6 + offset
        if row > 8:
            break
        last_row = row
        label = next(iter(periods.values()))[0].label
        sheet.cell(row, 1).value = label
        factor = (sign_rules or {}).get(key, 1.0)
        for year, items in periods.items():
            column = year_columns.get(year)
            if column:
                _set_income_value(sheet, sheet.cell(row, column), sum(item.value for item in items) * factor, items, mappings, label, "附注直取+符号规则" if factor != 1 else "附注直取")
    all_evidence = [item for periods in grouped.values() for items in periods.values() for item in items]
    if last_row < 8 and all_evidence:
        for year, column in year_columns.items():
            main_column = {2023: 3, 2024: 4, 2025: 5}.get(year)
            main_value = income["历利表"].cell(statement_row, main_column).value if main_column else None
            if not isinstance(main_value, (int, float)):
                continue
            written = sum(float(sheet.cell(row, column).value or 0) for row in range(6, last_row + 1))
            difference = float(main_value) - written
            if abs(difference) <= 0.000001:
                continue
            reconciliation_row = last_row + 1
            sheet.cell(reconciliation_row, 1).value = "主表与附注差异（待审计复核）"
            cell = sheet.cell(reconciliation_row, column); cell.value = difference
            _append_mapping(
                mappings, sheet, cell, sheet.cell(reconciliation_row, 1).value,
                all_evidence, "审计主表减附注明细", "差额保留并显式标记，不改写任何来源金额", "待复核",
            )


def _tax_rates(income, notes: list[NoteTable], mappings: list[list[Any]]) -> None:
    # Tax-rate tables sit before the consolidated/parent note sections, so
    # identify them by their row labels instead of scope/topic inheritance.
    for note in notes:
        matrix = note.table.matrix
        if not matrix:
            continue
        for row_number, row in enumerate(matrix, 1):
            label = str(row[0] if row else "").strip()
            rate_text = " ".join(str(value or "") for value in row[1:])
            match = re.search(r"(\d+(?:\.\d+)?)%", rate_text)
            if not match:
                continue
            rate = float(match.group(1)) / 100.0
            targets: list[tuple[str, str]] = []
            if label == "城市维护建设税": targets.append(("税金及附加表", "B6"))
            elif label == "教育费附加": targets.append(("税金及附加表", "B7"))
            elif label == "地方教育费附加": targets.append(("税金及附加表", "B8"))
            elif label == "企业所得税": targets.extend(("所得税表", f"{column}42") for column in "DEF")
            for sheet_name, coordinate in targets:
                sheet = income[sheet_name]; cell = sheet[coordinate]; cell.value = rate
                evidence = [NoteValue(label, 0, rate, note.table, row_number, 2)]
                _append_mapping(mappings, sheet, cell, label, evidence, "税率直取", "百分数文本转为Excel比例")


def _inventory_asset(asset, notes: list[NoteTable], mappings, scope: str) -> None:
    candidates = _selected(notes, "存货", scope)
    if not candidates:
        return
    current_year = 2025
    rows: dict[str, dict[str, list[NoteValue]]] = {}
    for note in candidates:
        matrix = note.table.matrix
        years = _year_columns(matrix)
        measures: dict[int, str] = {}
        if len(matrix) > 1:
            carried = ""
            for column, raw in enumerate(matrix[1], 1):
                text = str(raw or "").strip()
                if text: carried = text
                measures[column] = carried
        for row_number, row in enumerate(matrix[2:], 3):
            label = str(row[0] if row else "").strip()
            if not label: continue
            for column, year in years.items():
                if year != current_year: continue
                value = _ocr_number(row[column - 1] if len(row) >= column else "")
                if isinstance(value, (int, float)):
                    rows.setdefault(_compact(label), {}).setdefault(_compact(measures.get(column, "")), []).append(NoteValue(label, year, float(value), note.table, row_number, column))
    sheet = asset["表3-9"]
    target_rows = {"原材料": 9, "库存商品": 12, "半成品在产品": 13}
    for source_keys, target_key in [(("原材料",), "原材料"), (("库存商品",), "库存商品"), (("半成品", "在产品"), "半成品在产品")]:
        evidence: list[NoteValue] = []
        for source_key in source_keys:
            measures = rows.get(source_key, {})
            evidence.extend(measures.get("账面价值", []))
        if evidence:
            cell = sheet.cell(target_rows[target_key], 4); cell.value = sum(item.value for item in evidence)
            _append_mapping(mappings, sheet, cell, sheet.cell(target_rows[target_key], 3).value, evidence, "附注分类汇总", "母公司存货账面价值；人民币元直写")
    total_measures = rows.get("合计", {})
    for measure, target_row in (("账面余额", 16), ("存货跌价准备", 17), ("账面价值", 18)):
        evidence = total_measures.get(_compact(measure), [])
        if evidence:
            cell = sheet.cell(target_row, 4); cell.value = sum(item.value for item in evidence)
            _append_mapping(mappings, sheet, cell, sheet.cell(target_row, 3).value, evidence, "附注合计直取", "人民币元直写；母公司口径")


def _fixed_assets(asset, income, notes, asset_mappings, income_mappings, scope: str) -> None:
    selected = sorted(_selected(notes, "固定资产", scope), key=lambda item: item.table.page_number)
    if not selected:
        return
    section = ""
    original: dict[str, list[NoteValue]] = {}
    net: dict[str, list[NoteValue]] = {}
    original_total: list[NoteValue] = []
    net_total: list[NoteValue] = []
    for note in selected:
        matrix = note.table.matrix
        years = _year_columns(matrix)
        for row_number, row in enumerate(matrix, 1):
            label = str(row[0] if row else "").strip()
            key = _compact(label)
            if "账面原值合计" in key:
                section = "original"
            elif "累计折旧合计" in key:
                section = "depreciation"
            elif "账面净值合计" in key or "账面价值合计" in key:
                section = "net"
            elif re.match(r"^[一二三四五六七八九十]+、", label):
                section = "other"
            for column, year in years.items():
                if year != 2025:
                    continue
                value = _ocr_number(row[column - 1] if len(row) >= column else "")
                if not isinstance(value, (int, float)):
                    continue
                evidence = NoteValue(label, year, float(value), note.table, row_number, column)
                if "账面原值合计" in key:
                    original_total = [evidence]
                elif "账面净值合计" in key or "账面价值合计" in key:
                    net_total = [evidence]
                elif re.match(r"^其中", label):
                    category = _compact(re.sub(r"^其中[：:]?", "", label))
                    if section == "original": original.setdefault(category, []).append(evidence)
                    elif section == "net": net.setdefault(category, []).append(evidence)
                elif key in {"生产设备", "办公设备", "装修工程", "研发设备", "工装模具", "运输设备", "电子设备"}:
                    if section == "original": original.setdefault(key, []).append(evidence)
                    elif section == "net": net.setdefault(key, []).append(evidence)

    groups = {
        "房屋建筑物": ("房屋建筑物",),
        "机器设备": ("生产设备", "研发设备", "工装模具"),
        "运输设备": ("运输设备",),
        "电子设备和其他设备": ("办公设备", "电子设备"),
        "其他资产": ("装修工程",),
    }
    asset_sheet = asset["表4-6"]
    asset_rows = {"房屋建筑物": 9, "机器设备": 13, "运输设备": 14, "电子设备和其他设备": 15}
    def deduplicate(items: list[NoteValue]) -> list[NoteValue]:
        unique: dict[float, NoteValue] = {}
        for item in items:
            unique.setdefault(round(item.value, 6), item)
        return list(unique.values())
    for group, source_keys in groups.items():
        original_evidence = [item for key in source_keys for item in deduplicate(original.get(key, []))]
        net_evidence = [item for key in source_keys for item in deduplicate(net.get(key, []))]
        if group in asset_rows:
            row = asset_rows[group]
            if original_evidence:
                cell = asset_sheet.cell(row, 4); cell.value = sum(item.value for item in original_evidence)
                _append_mapping(asset_mappings, asset_sheet, cell, asset_sheet.cell(row, 3).value, original_evidence, "固定资产附注分类汇总", "人民币元；母公司口径；类别别名归一化")
            if net_evidence:
                cell = asset_sheet.cell(row, 5); cell.value = sum(item.value for item in net_evidence)
                _append_mapping(asset_mappings, asset_sheet, cell, asset_sheet.cell(row, 3).value, net_evidence, "固定资产附注分类汇总", "人民币元；母公司口径；类别别名归一化")
        income_row = {"房屋建筑物": 7, "机器设备": 11, "运输设备": 15, "电子设备和其他设备": 19, "其他资产": 27}[group]
        income_sheet = income["折旧摊销附表"]
        if original_evidence or net_evidence:
            income_sheet.cell(income_row, 2).value = "审计附注分类合计"
        if original_evidence:
            cell = income_sheet.cell(income_row, 5); cell.value = sum(item.value for item in original_evidence) / 10_000.0
            _append_mapping(income_mappings, income_sheet, cell, f"{group}账面原值", original_evidence, "固定资产附注分类汇总", "人民币元÷10,000；只填基准日原值，不推断寿命/残值率")
        if net_evidence:
            cell = income_sheet.cell(income_row, 6); cell.value = sum(item.value for item in net_evidence) / 10_000.0
            _append_mapping(income_mappings, income_sheet, cell, f"{group}账面净值", net_evidence, "固定资产附注分类汇总", "人民币元÷10,000；只填基准日净值，不推断寿命/残值率")
    if original_total:
        cell = asset_sheet.cell(17, 4); cell.value = original_total[-1].value
        _append_mapping(asset_mappings, asset_sheet, cell, "固定资产账面原值合计", original_total, "附注合计直取", "人民币元；母公司口径")
    if net_total:
        cell = asset_sheet.cell(19, 5); cell.value = net_total[-1].value
        _append_mapping(asset_mappings, asset_sheet, cell, "固定资产账面价值合计", net_total, "附注合计直取", "人民币元；母公司口径")


def _bank_deposit_asset(asset, notes, mappings, scope: str) -> None:
    values = [value for note in _selected(notes, "货币资金", scope) for value in _simple_note_values(note)]
    current = [value for value in values if value.period == 2025 and _compact(value.label) == "银行存款"]
    if not current: return
    sheet = asset["表3-1"]
    for row in (9, 11):
        cell = sheet.cell(row, 4); cell.value = sum(value.value for value in current)
        _append_mapping(mappings, sheet, cell, sheet.cell(row, 3).value, current, "附注直取/汇总", "人民币元直写；审计附注仅披露银行存款合计")


def _employee_tax_asset(asset, notes, mappings, scope: str) -> None:
    configs = [
        ("应付职工薪酬", "表5-6", 2, 4, 48),
        ("应交税费", "表5-7", 4, 5, 48),
    ]
    for topic, sheet_name, label_column, value_column, total_row in configs:
        values = [value for note in _selected(notes, topic, scope) for value in _simple_note_values(note)]
        grouped: dict[str, list[NoteValue]] = {}
        for value in values:
            key = _compact(value.label)
            if value.period == 2025 and "合计" not in key and "总计" not in key:
                grouped.setdefault(key, []).append(value)
        sheet = asset[sheet_name]
        for offset, items in enumerate(grouped.values()):
            row = 8 + offset
            if row >= total_row: break
            sheet.cell(row, 1).value = offset + 1
            sheet.cell(row, label_column).value = items[0].label
            cell = sheet.cell(row, value_column); cell.value = sum(item.value for item in items)
            _append_mapping(mappings, sheet, cell, items[0].label, items, "附注明细直取", "人民币元直写；发生日期/结算对象未披露则留空")


def _long_term_investments(asset, income, notes, mappings_asset, mappings_income, scope: str) -> None:
    selected = _selected(notes, "长期股权投资", scope)
    values = [value for note in selected for value in _simple_note_values(note)]
    grouped: dict[str, list[NoteValue]] = {}
    for value in values:
        key = _compact(value.label)
        if value.period == 2025 and "合计" not in key and "小计" not in key and "控制企业" not in key:
            grouped.setdefault(key, []).append(value)
    asset_sheet = asset["表4-4"]
    income_sheet = income["长期股权投资表"]
    # The sample's hard-coded accounting method belongs to Tongfu and is not
    # a structural constant.
    for row in range(5, 55):
        if income_sheet.cell(row, 2).value == "权益法": income_sheet.cell(row, 2).value = None
    for offset, items in enumerate(grouped.values()):
        asset_row = 8 + offset; income_row = 5 + offset
        label = items[0].label; amount = sum(item.value for item in items)
        asset_sheet.cell(asset_row, 1).value = offset + 1
        asset_sheet.cell(asset_row, 2).value = label
        asset_cell = asset_sheet.cell(asset_row, 7); asset_cell.value = amount
        _append_mapping(mappings_asset, asset_sheet, asset_cell, label, items, "附注明细直取", "母公司长期股权投资2025年末账面价值；人民币元")
        income_sheet.cell(income_row, 1).value = label
        income_cell = income_sheet.cell(income_row, 3); income_cell.value = amount / 10_000.0
        _append_mapping(mappings_income, income_sheet, income_cell, label, items, "附注明细直取", "人民币元÷10,000写入万元；核算方法无直接证据留空")


def _interest_debt_income(income, notes, mappings, scope: str) -> None:
    sheet = income["付息负债明细表"]
    entries: list[tuple[str, list[NoteValue]]] = []
    for topic, label in (("一年内到期的非流动负债", "一年内到期的长期借款"), ("长期借款", "保证借款")):
        values = [value for note in _selected(notes, topic, scope) for value in _simple_note_values(note)]
        matched = [value for value in values if value.period == 2025 and _compact(value.label) == _compact(label)]
        if matched: entries.append((label, matched))
    for offset, (label, items) in enumerate(entries):
        row = 6 + offset
        sheet.cell(row, 1).value = "银行借款"
        sheet.cell(row, 2).value = label
        cell = sheet.cell(row, 7); cell.value = sum(item.value for item in items) / 10_000.0
        sheet.cell(row, 9).value = "审计附注仅披露期末余额；合同日、期限、利率、用途待借款合同"
        _append_mapping(mappings, sheet, cell, label, items, "附注直取", "人民币元÷10,000写入万元；未披露字段留空")


def _main_products(income, notes, mappings, scope: str) -> None:
    selected = _selected(notes, "营业收入和营业成本", scope)
    if not selected: return
    sheet = income["主要产品及服务"]
    labels = [(5, "主营业务", "主营业务（审计附注未进一步按产品拆分）"), (9, "其他业务", "其他业务")]
    for sequence, (target_row, source_label, display) in enumerate(labels, 1):
        evidence: list[NoteValue] = []
        for note in selected:
            for row_number, row in enumerate(note.table.matrix, 1):
                if row and str(row[0]).strip() == source_label:
                    evidence.append(NoteValue(source_label, 0, 0.0, note.table, row_number, 1))
        if not evidence: continue
        sheet.cell(target_row, 1).value = source_label
        sheet.cell(target_row, 2).value = sequence
        cell = sheet.cell(target_row, 3); cell.value = display
        sheet.cell(target_row, 4).value = None
        sheet.cell(target_row, 5).value = "待管理层预测资料"
        sheet.cell(target_row, 6).value = "待管理层预测资料"
        sheet.cell(target_row, 7).value = "仅据审计附注分类；产品、销量、单价及预测方法待管理层资料"
        _append_mapping(mappings, sheet, cell, source_label, evidence, "附注业务分类直取", "不推断具体产品或预测方法")


def enhance_from_notes(
    *, asset, income, document_name: str, pages: list[dict[str, Any]],
    asset_mappings: list[list[Any]], income_mappings: list[list[Any]], scope: str = "母公司",
    statement_source_rows: dict[str, list[Any]] | None = None,
) -> list[NoteTable]:
    """Populate all evidence-backed note detail while preserving formulas."""
    notes = note_tables_from_pages(document_name, pages)
    _income_statement_split(income, notes, income_mappings, scope)
    note_totals = _fill_income_totals_from_notes(income, notes, income_mappings, scope)
    if statement_source_rows:
        _reconcile_operating_profit(income, note_totals, income_mappings, statement_source_rows, scope)
    _main_products(income, notes, income_mappings, scope)
    for topic, sheet_name, total_row, statement_row in (
        ("销售费用", "销售费用表", 62, 25),
        ("管理费用", "管理费用表", 62, 26),
        ("研发费用", "研发费用表", 63, 27),
    ):
        _expense_details(income, notes, income_mappings, topic, sheet_name, total_row, scope, statement_row)
    _expense_details(income, notes, income_mappings, "财务费用", "财务费用表", 15, scope, 28)
    for topic, sheet_name, statement_row in (
        ("其他收益", "其他收益表", 31),
        ("信用减值损失", "信用减值损失表", 38),
        ("资产减值损失", "资产减值损失表", 39),
        ("资产处置收益", "资产处置收益表", 40),
        ("营业外收入", "营业外收入表", 42),
        ("营业外支出", "营业外支出表", 43),
    ):
        _small_detail_sheet(income, notes, income_mappings, topic, sheet_name, scope, statement_row)
    _tax_rates(income, notes, income_mappings)
    _bank_deposit_asset(asset, notes, asset_mappings, scope)
    _inventory_asset(asset, notes, asset_mappings, scope)
    _fixed_assets(asset, income, notes, asset_mappings, income_mappings, scope)
    _employee_tax_asset(asset, notes, asset_mappings, scope)
    _long_term_investments(asset, income, notes, asset_mappings, income_mappings, scope)
    _interest_debt_income(income, notes, income_mappings, scope)
    return notes
