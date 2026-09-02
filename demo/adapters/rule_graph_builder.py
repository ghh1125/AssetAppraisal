"""Build a machine-readable, per-cell appraisal mapping and formula graph.

The graph deliberately separates reusable template logic from case evidence:
formulas and blank input requirements come from the Tongfu-derived template,
while populated evidence nodes come only from the current archive's trace
ledger.  It is therefore safe to reuse for another company.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from hashlib import sha1
import json
from pathlib import Path
import re
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter, range_boundaries


TRACE_HEADERS = (
    "目标文件", "目标工作表", "目标单元格", "模板字段", "来源文件", "来源页",
    "来源表", "来源行", "取数方式", "计算/处理", "状态",
)


def _node_id(*parts: str) -> str:
    return "n_" + sha1("|".join(parts).encode("utf-8")).hexdigest()[:14]


def _mmd(value: Any, limit: int = 80) -> str:
    text = str(value if value is not None else "").replace("\n", " ").replace('"', "'")
    return text if len(text) <= limit else text[: limit - 1] + "…"


def _safe_filename(value: str) -> str:
    return re.sub(r"[\\/*?:\[\]]", "_", value)


def _field_label(sheet, row: int, column: int) -> str:
    for candidate in range(column - 1, max(0, column - 6), -1):
        value = sheet.cell(row, candidate).value
        if isinstance(value, str) and value and not value.startswith("="):
            return value.strip()
    for candidate in range(1, min(sheet.max_column, 8) + 1):
        value = sheet.cell(row, candidate).value
        if isinstance(value, str) and value and not value.startswith("="):
            return value.strip()
    return ""


def _trace_rows(trace_path: Path) -> tuple[dict[tuple[str, str, str], list[dict[str, Any]]], list[dict[str, Any]]]:
    workbook = load_workbook(trace_path, read_only=True, data_only=True)
    sheet = workbook["AI映射规则"]
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value or "") for value in rows[0]] if rows else list(TRACE_HEADERS)
    indexed: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    flat: list[dict[str, Any]] = []
    for values in rows[1:]:
        item = {headers[index]: values[index] if index < len(values) else None for index in range(len(headers))}
        key = (str(item.get("目标文件") or ""), str(item.get("目标工作表") or ""), str(item.get("目标单元格") or ""))
        indexed[key].append(item)
        flat.append(item)
    workbook.close()
    return indexed, flat


def _split_reference(token: str, current_sheet: str) -> tuple[str, str] | None:
    token = token.strip()
    if "[" in token or "]" in token:
        return None
    if "!" in token:
        sheet_name, address = token.rsplit("!", 1)
        sheet_name = sheet_name.strip("'").replace("''", "'")
    else:
        sheet_name, address = current_sheet, token
    address = address.replace("$", "")
    if not re.fullmatch(r"[A-Z]{1,3}\d+(?::[A-Z]{1,3}\d+)?", address, re.I):
        return None
    return sheet_name, address.upper()


def _formula_precedents(formula: str, current_sheet: str, limit: int = 5000) -> list[tuple[str, str]]:
    precedents: list[tuple[str, str]] = []
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return precedents
    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        parsed = _split_reference(token.value, current_sheet)
        if parsed is None:
            continue
        sheet_name, address = parsed
        if ":" not in address:
            precedents.append((sheet_name, address))
            continue
        min_col, min_row, max_col, max_row = range_boundaries(address)
        size = (max_col - min_col + 1) * (max_row - min_row + 1)
        if size > limit:
            continue
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                precedents.append((sheet_name, f"{get_column_letter(column)}{row}"))
    return list(dict.fromkeys(precedents))


def _role(value: Any, mappings: list[dict[str, Any]]) -> str:
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    if mappings:
        return "evidence_input" if any(item.get("状态") == "已填" for item in mappings) else "required_missing"
    if isinstance(value, (int, float)):
        return "template_constant"
    return "structural_label"


def build_rule_graph_bundle(
    *, output_dir: Path, workbook_paths: dict[str, Path], trace_path: Path,
    case_name: str = "当前案例",
) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    indexed_trace, flat_trace = _trace_rows(trace_path)
    nodes: dict[str, dict[str, Any]] = {}
    edges: list[dict[str, Any]] = []
    sheet_stats: list[dict[str, Any]] = []
    book_sheet_nodes: dict[tuple[str, str], list[str]] = defaultdict(list)
    dependency_counts: Counter[tuple[str, str, str, str]] = Counter()

    for workbook_name, workbook_path in workbook_paths.items():
        workbook = load_workbook(workbook_path, data_only=False, read_only=False)
        for sheet in workbook.worksheets:
            mapping_cells = {
                key[2] for key in indexed_trace
                if key[0] == workbook_name and key[1] == sheet.title
            }
            coordinates = {
                cell.coordinate
                for cell in sheet._cells.values()
                if cell.value not in (None, "")
            } | mapping_cells
            counts: Counter[str] = Counter()
            for coordinate in sorted(coordinates, key=lambda c: (sheet[c].row, sheet[c].column)):
                cell = sheet[coordinate]
                mappings = indexed_trace.get((workbook_name, sheet.title, coordinate), [])
                role = _role(cell.value, mappings)
                counts[role] += 1
                node_id = _node_id(workbook_name, sheet.title, coordinate)
                node = {
                    "id": node_id,
                    "workbook": workbook_name,
                    "sheet": sheet.title,
                    "cell": coordinate,
                    "role": role,
                    "field": (mappings[0].get("模板字段") if mappings else None) or _field_label(sheet, cell.row, cell.column),
                    "value_or_formula": cell.value,
                    "number_format": cell.number_format,
                    "mappings": mappings,
                    "generic_rule": (
                        ("保留模板原计算逻辑；当前缺资料导致显示错误时仅外包IFERROR并返回空白" if str(cell.value).upper().startswith("=IFERROR(") else "保留模板公式并由前置单元格驱动") if role == "formula" else
                        "仅用主体、口径、期间、单位均匹配的来源证据填入" if role == "evidence_input" else
                        "材料无充分依据则留空，并按缺失资料要求补件" if role == "required_missing" else
                        "模板结构/标题/常量；换公司时不得作为事实来源"
                    ),
                }
                nodes[node_id] = node
                book_sheet_nodes[(workbook_name, sheet.title)].append(node_id)
                if role == "formula":
                    for source_sheet, source_cell in _formula_precedents(str(cell.value), sheet.title):
                        source_id = _node_id(workbook_name, source_sheet, source_cell)
                        if source_id not in nodes:
                            nodes[source_id] = {
                                "id": source_id, "workbook": workbook_name, "sheet": source_sheet,
                                "cell": source_cell, "role": "formula_precedent", "field": "",
                                "value_or_formula": None, "number_format": "", "mappings": [],
                                "generic_rule": "由被引用单元格提供前置值",
                            }
                        edges.append({"from": source_id, "to": node_id, "type": "formula_reference"})
                        dependency_counts[(workbook_name, source_sheet, workbook_name, sheet.title)] += 1
                for mapping in mappings:
                    if mapping.get("来源文件") and mapping.get("来源页"):
                        evidence_key = f"{mapping.get('来源文件')}|p{mapping.get('来源页')}|{mapping.get('来源表')}|{mapping.get('来源行')}"
                        evidence_id = _node_id("evidence", evidence_key)
                        nodes.setdefault(evidence_id, {
                            "id": evidence_id, "role": "source_evidence", "source_file": mapping.get("来源文件"),
                            "page": mapping.get("来源页"), "table": mapping.get("来源表"), "row": mapping.get("来源行"),
                            "generic_rule": "OCR原文定位；复核时可回到文件/页/表/行",
                        })
                        edges.append({"from": evidence_id, "to": node_id, "type": "evidence_mapping", "method": mapping.get("取数方式"), "transform": mapping.get("计算/处理")})
            sheet_stats.append({"workbook": workbook_name, "sheet": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column, **counts})
        workbook.close()

    graph = {
        "version": "appraisal_cell_graph.v3",
        "case": case_name,
        "principles": [
            "模板只提供Sheet结构、单元格角色和公式链，不提供当前案例数值",
            "RAR材料先做主体匹配，再按母公司/合并口径、期间和单位匹配",
            "审计主表优先；附注用于分类拆分、明细和交叉复核",
            "直接取数必须保留文件、页、表、行列；换算必须记录公式",
            "评估值、设备参数、预测、WACC等无证据字段必须留空并列补件要求",
        ],
        "workbooks": {name: str(path) for name, path in workbook_paths.items()},
        "trace_workbook": str(trace_path),
        "statistics": {
            "sheet_count": len(sheet_stats), "node_count": len(nodes), "edge_count": len(edges),
            "formula_node_count": sum(item.get("formula", 0) for item in sheet_stats),
            "evidence_input_count": sum(item.get("evidence_input", 0) for item in sheet_stats),
            "required_missing_count": sum(item.get("required_missing", 0) for item in sheet_stats),
        },
        "sheet_statistics": sheet_stats,
        "nodes": list(nodes.values()),
        "edges": edges,
    }
    (output_dir / "资产评估通用逐单元格规则图.json").write_text(json.dumps(graph, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    # A zero-dependency visual browser makes the complete graph inspectable
    # without requiring Mermaid/Graphviz to be installed.  Users can select a
    # workbook and Sheet, search by cell/field/source, and inspect each edge.
    browser_payload = json.dumps(graph, ensure_ascii=False, separators=(",", ":"), default=str).replace("</", "<\\/")
    browser = r'''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>资产评估逐单元格规则图</title>
<style>body{font-family:"Microsoft YaHei",sans-serif;margin:0;background:#f4f7fb;color:#182230}header{padding:18px 24px;background:linear-gradient(135deg,#15395f,#255a93);color:white}header h1{margin:0 0 6px;font-size:22px}.bar{display:flex;gap:10px;align-items:center;flex-wrap:wrap;padding:14px 24px;background:white;position:sticky;top:0;z-index:2;box-shadow:0 2px 12px #19365d14}select,input{padding:8px 10px;border:1px solid #cbd6e3;border-radius:8px;min-width:210px;background:#fbfdff}.stats{margin-left:auto;color:#56677b;font-size:12px}.legend{display:flex;gap:8px;flex-wrap:wrap;padding:10px 24px 0;color:#52667b;font-size:12px}.legend span{display:inline-flex;align-items:center;gap:5px;padding:4px 7px;border-radius:999px;background:#fff;box-shadow:0 1px 4px #19365d12}.legend i{width:10px;height:10px;border-radius:50%;display:inline-block}.main{display:grid;grid-template-columns:minmax(620px,2fr) minmax(390px,1fr);gap:14px;padding:14px 24px}.panel{background:white;border-radius:12px;box-shadow:0 4px 14px #19365d12;padding:12px;overflow:auto}svg{width:100%;min-height:680px;background:linear-gradient(180deg,#fcfdff,#f7faff);border-radius:9px}.node{cursor:pointer}.node rect{stroke-width:1.4;filter:drop-shadow(0 2px 2px #19365d20)}.node text{font-size:11px;fill:#182230}.node .node-role{font-size:10px;fill:#52667b}.node:hover rect{stroke-width:2.6}.edge{stroke:#9cadc0;stroke-width:1.15;stroke-opacity:.76;fill:none;marker-end:url(#arrow)}.edge.evidence_mapping{stroke:#6bae7a;stroke-dasharray:4 3}.formula rect{fill:#e6f0ff;stroke:#4e83d8}.evidence_input rect{fill:#e0f5e6;stroke:#4a9c68}.required_missing rect{fill:#fff0df;stroke:#e1842d}.source_evidence rect{fill:#fff8cf;stroke:#c99e22}.formula_precedent rect{fill:#eee9ff;stroke:#8169c7}.template_constant rect{fill:#e9f0f5;stroke:#7890a5}table{border-collapse:collapse;width:100%;font-size:12px}th,td{border-bottom:1px solid #e4e9ef;padding:7px;text-align:left;vertical-align:top}tr{cursor:pointer}tr:hover{background:#f4f8ff}th{background:#eef3f8;position:sticky;top:0}.detail pre{white-space:pre-wrap;word-break:break-word;font-size:12px}.role-badge{display:inline-block;padding:2px 6px;border-radius:99px;font-size:11px}.role-formula{background:#e6f0ff;color:#245da8}.role-evidence_input{background:#e0f5e6;color:#287145}.role-required_missing{background:#fff0df;color:#a9550c}.role-source_evidence{background:#fff8cf;color:#80610a}.role-formula_precedent{background:#eee9ff;color:#5e4ba4}@media(max-width:1000px){.main{grid-template-columns:1fr}.stats{margin-left:0}.legend{padding-left:14px}.main{padding:14px}}</style></head>
<body><header><h1>资产评估逐单元格规则图</h1><div>模板公式链 + 当前案例证据定位 + 缺失资料要求；模板数值不作为案例事实。</div></header>
<div class="bar"><select id="book"></select><select id="sheet"></select><select id="mode"><option value="rules">重点规则节点</option><option value="all">全部有值单元格</option></select><input id="q" placeholder="搜索单元格、字段、来源或公式"><span class="stats" id="stats"></span></div><div class="legend"><span><i style="background:#4e83d8"></i>模板公式</span><span><i style="background:#4a9c68"></i>已取证填入</span><span><i style="background:#e1842d"></i>待补资料</span><span><i style="background:#c99e22"></i>来源证据</span><span><i style="background:#8169c7"></i>公式前置</span></div>
<div class="main"><div class="panel"><svg id="graph"><defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="3" orient="auto"><path d="M0,0 L0,6 L7,3 z" fill="#93a1b2"/></marker></defs></svg></div><div class="panel detail"><h3>单元格规则明细</h3><div id="detail">点击图中节点或右侧表格行。</div><table><thead><tr><th>单元格</th><th>角色</th><th>字段/规则</th></tr></thead><tbody id="rows"></tbody></table></div></div>
<script>const G=__GRAPH__;const nodes=new Map(G.nodes.map(n=>[n.id,n]));const edges=G.edges;const book=document.querySelector('#book'),sheet=document.querySelector('#sheet'),mode=document.querySelector('#mode'),q=document.querySelector('#q');
const books=[...new Set([...nodes.values()].map(n=>n.workbook))];books.forEach(x=>book.add(new Option(x,x)));function sheets(){sheet.innerHTML='';[...new Set([...nodes.values()].filter(n=>n.workbook===book.value).map(n=>n.sheet))].forEach(x=>sheet.add(new Option(x,x)));render()}
function esc(x){return String(x??'').replace(/[&<>]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;'}[c]))}function title(n){return n.cell?(n.cell+' '+String(n.field||'').slice(0,23)):(String(n.source_file||'来源证据').slice(0,25)+' p'+String(n.page||''))}function show(n){const location=n.cell?(esc(n.workbook)+' / '+esc(n.sheet)+'!'+esc(n.cell)):esc(n.source_file)+' 第'+esc(n.page)+'页';document.querySelector('#detail').innerHTML='<b>'+location+'</b><pre>'+esc(JSON.stringify(n,null,2))+'</pre>'}
function render(){const term=q.value.trim().toLowerCase(),focus=[...nodes.values()].filter(n=>n.workbook===book.value&&n.sheet===sheet.value&&n.cell);let all=focus.filter(n=>mode.value==='all'||['formula','evidence_input','required_missing','template_constant'].includes(n.role));if(term)all=all.filter(n=>JSON.stringify(n).toLowerCase().includes(term));const focusIds=new Set(all.map(n=>n.id)),related=new Set();edges.forEach(e=>{if(focusIds.has(e.to)&&nodes.has(e.from))related.add(e.from);if(focusIds.has(e.from)&&nodes.has(e.to))related.add(e.to)});let shown=[...all,...[...related].map(id=>nodes.get(id)).filter(n=>n&&!focusIds.has(n.id))];if(term)shown=shown.filter(n=>JSON.stringify(n).toLowerCase().includes(term)||focusIds.has(n.id));shown=shown.slice(0,260);const ids=new Set(shown.map(n=>n.id));document.querySelector('#stats').textContent='本 Sheet 展示 '+shown.length+' 节点；全库 '+G.statistics.node_count+' 节点 / '+G.statistics.edge_count+' 边';
const cols={source_evidence:0,evidence_input:1,required_missing:2,formula:3,formula_precedent:4,template_constant:4},ys=[0,0,0,0,0],pos=new Map();shown.forEach(n=>{const c=cols[n.role]??4;pos.set(n.id,{x:18+c*244,y:20+ys[c]*64});ys[c]++});const svg=document.querySelector('#graph');svg.setAttribute('viewBox','0 0 1235 '+Math.max(680,Math.max(...ys)*64+60));svg.querySelectorAll('.dynamic').forEach(x=>x.remove());
edges.filter(e=>ids.has(e.from)&&ids.has(e.to)).forEach(e=>{const a=pos.get(e.from),b=pos.get(e.to);if(!a||!b)return;const l=document.createElementNS('http://www.w3.org/2000/svg','line');l.setAttribute('x1',a.x+220);l.setAttribute('y1',a.y+21);l.setAttribute('x2',b.x);l.setAttribute('y2',b.y+21);l.setAttribute('class','dynamic edge '+e.type);svg.append(l)});shown.forEach(n=>{const p=pos.get(n.id),g=document.createElementNS('http://www.w3.org/2000/svg','g');g.setAttribute('class','dynamic node '+n.role);g.setAttribute('transform','translate('+p.x+','+p.y+')');const r=document.createElementNS('http://www.w3.org/2000/svg','rect');r.setAttribute('width',220);r.setAttribute('height',46);r.setAttribute('rx',10);const t=document.createElementNS('http://www.w3.org/2000/svg','text');t.setAttribute('x',9);t.setAttribute('y',18);t.textContent=title(n);const t2=document.createElementNS('http://www.w3.org/2000/svg','text');t2.setAttribute('class','node-role');t2.setAttribute('x',9);t2.setAttribute('y',35);t2.textContent=n.role;g.append(r,t,t2);g.onclick=()=>show(n);svg.append(g)});
const tbody=document.querySelector('#rows');tbody.innerHTML='';all.forEach(n=>{const tr=document.createElement('tr');tr.innerHTML='<td>'+esc(n.cell)+'</td><td><span class="role-badge role-'+esc(n.role)+'">'+esc(n.role)+'</span></td><td>'+esc(n.field||n.generic_rule)+'</td>';tr.onclick=()=>show(n);tbody.append(tr)})}
book.onchange=sheets;sheet.onchange=render;mode.onchange=render;q.oninput=render;sheets();</script></body></html>'''.replace("__GRAPH__", browser_payload)
    (output_dir / "逐单元格规则图浏览器.html").write_text(browser, encoding="utf-8")

    missing = [item for item in flat_trace if item.get("状态") != "已填"]
    grouped_missing: dict[str, dict[str, list[dict[str, Any]]]] = defaultdict(lambda: defaultdict(list))
    for item in missing:
        grouped_missing[str(item.get("目标文件") or "")][str(item.get("目标工作表") or "")].append({
            "cell": item.get("目标单元格"), "field": item.get("模板字段"),
            "reason": item.get("计算/处理"), "required_material": (
                "预测/估值参数或业务明细资料" if any(token in str(item.get("模板字段") or "") for token in ("预测", "评估", "折现", "WACC", "设备", "房屋"))
                else "能证明该字段金额、期间、单位和主体口径的原始资料"
            ),
        })
    (output_dir / "当前案例待补充资料_逐单元格.json").write_text(json.dumps(grouped_missing, ensure_ascii=False, indent=2), encoding="utf-8")

    master = ["flowchart LR", '  archive["新RAR材料包"] --> identity["主体识别与材料归属"]', '  identity --> ocr["逐页OCR与表格坐标"]', '  ocr --> scope["口径/期间/单位判定"]']
    for workbook_name in workbook_paths:
        book_id = _node_id("book", workbook_name)
        master.append(f'  scope --> {book_id}["{_mmd(workbook_name)}"]')
        for stats in sheet_stats:
            if stats["workbook"] == workbook_name:
                sheet_id = _node_id("sheet", workbook_name, stats["sheet"])
                master.append(f'  {book_id} --> {sheet_id}["{_mmd(stats["sheet"])}"]')
    for (from_book, from_sheet, to_book, to_sheet), count in dependency_counts.items():
        source_id = _node_id("sheet", from_book, from_sheet)
        target_id = _node_id("sheet", to_book, to_sheet)
        master.append(f'  {source_id}["{_mmd(from_sheet)}"] -->|{count}个公式引用| {target_id}["{_mmd(to_sheet)}"]')
    (output_dir / "00_总流程与Sheet依赖.mmd").write_text("\n".join(dict.fromkeys(master)), encoding="utf-8")

    index_lines = ["# 逐单元格规则图索引", "", "每个 Sheet 单独一张图；JSON 是自动生成器实际读取的完整规则库。", ""]
    for sequence, stats in enumerate(sheet_stats, 1):
        workbook_name, sheet_name = stats["workbook"], stats["sheet"]
        filename = f"{sequence:02d}_{_safe_filename(workbook_name)}_{_safe_filename(sheet_name)}.mmd"
        lines = ["flowchart LR"]
        relevant = book_sheet_nodes[(workbook_name, sheet_name)]
        relevant_set = set(relevant)
        for node_id in relevant:
            node = nodes[node_id]
            if node["role"] not in {"formula", "evidence_input", "required_missing"}:
                continue
            shape = "{{" if node["role"] == "required_missing" else "["
            close = "}}" if node["role"] == "required_missing" else "]"
            lines.append(f'  {node_id}{shape}"{_mmd(node["cell"] + " " + str(node.get("field") or node.get("value_or_formula") or ""))}"{close}')
        for edge in edges:
            if edge["to"] not in relevant_set:
                continue
            if edge["type"] == "formula_reference" and edge["from"] in relevant_set:
                lines.append(f'  {edge["from"]} --> {edge["to"]}')
            elif edge["type"] == "evidence_mapping":
                evidence = nodes[edge["from"]]
                lines.append(f'  {edge["from"]}["p{_mmd(evidence.get("page"))} {_mmd(evidence.get("row"))}"] -->|{_mmd(edge.get("method"), 24)}| {edge["to"]}')
        (output_dir / filename).write_text("\n".join(dict.fromkeys(lines)), encoding="utf-8")
        index_lines.append(
            f"- `{workbook_name}` / `{sheet_name}` → `{filename}`；公式 {stats.get('formula', 0)}，证据输入 {stats.get('evidence_input', 0)}，待补 {stats.get('required_missing', 0)}"
        )
    (output_dir / "规则图索引.md").write_text("\n".join(index_lines), encoding="utf-8")
    return graph["statistics"]
